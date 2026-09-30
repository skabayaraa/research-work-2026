"""ШУТИС-ийн 2024 оны загвар дээр өгүүллийг угсрах скрипт."""
import copy
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

W = Path("/home/claude/work")
PRJ = W / "project"
FIG = PRJ / "paper_figs"
R = PRJ / "ml" / "results"

doc = Document(str(W / "template.docx"))
body = doc.element.body
els = list(body)

# ------------------------------------------------------------------ тоон үр дүн
df = pd.read_csv(R / "all_results.csv")
g = df.groupby("model")
M = lambda m, c: g[c].mean()[m]
S = lambda m, c: g[c].std()[m]
ex = json.load(open(R / "extras.json"))
lat = json.load(open(R / "latency.json"))
nlpc = json.load(open(R / "nlp_compare.json"))
T = ex["tests"]


def f3(x):
    return f"{x:.3f}"


def pm(m, c):
    return f"{M(m, c):.3f}±{S(m, c):.3f}"


# ------------------------------------------------------------------ туслах
def set_runs(p_el, parts, base_rpr=None):
    """p_el доторх run-уудыг устгаж, шинэ run нэмнэ. parts: [(text, {b,i,sup,sub,color})]."""
    for r in list(p_el):
        if r.tag in (qn("w:r"), qn("w:proofErr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:hyperlink")):
            p_el.remove(r)
    for text, fmt in parts:
        r = OxmlElement("w:r")
        rpr = copy.deepcopy(base_rpr) if base_rpr is not None else OxmlElement("w:rPr")
        for tag in ("w:vertAlign",):
            for e in rpr.findall(qn(tag)):
                rpr.remove(e)
        if fmt.get("sup") or fmt.get("sub"):
            va = OxmlElement("w:vertAlign")
            va.set(qn("w:val"), "superscript" if fmt.get("sup") else "subscript")
            rpr.append(va)
        if fmt.get("b"):
            rpr.insert(0, OxmlElement("w:b"))
        if fmt.get("i"):
            rpr.insert(0, OxmlElement("w:i"))
        r.append(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p_el.append(r)


def first_rpr(p_el):
    r = p_el.find(qn("w:r"))
    rp = r.find(qn("w:rPr")) if r is not None else None
    return copy.deepcopy(rp) if rp is not None else OxmlElement("w:rPr")


def parse_inline(s):
    """'**тод**', '_налуу_', '^дээд^', '~доод~' тэмдэглэгээг run болгон задлах."""
    import re
    out = []
    for tok in re.split(r"(\*\*.+?\*\*|_[^_]+?_|\^[^^]+?\^|~[^~]+?~)", s):
        if not tok:
            continue
        if tok.startswith("**"):
            out.append((tok[2:-2], {"b": True}))
        elif tok.startswith("_") and tok.endswith("_") and len(tok) > 2:
            out.append((tok[1:-1], {"i": True}))
        elif tok.startswith("^"):
            out.append((tok[1:-1], {"sup": True}))
        elif tok.startswith("~"):
            out.append((tok[1:-1], {"sub": True}))
        else:
            out.append((tok, {}))
    return out


def add_p(text, style="Para 1", align=None, keep_next=False, space_after=None):
    p = doc.add_paragraph(style=style)
    for t, f in parse_inline(text):
        run = p.add_run(t)
        run.bold = True if f.get("b") else None
        run.italic = True if f.get("i") else None
        if f.get("sup"):
            run.font.superscript = True
        if f.get("sub"):
            run.font.subscript = True
    if align:
        p.alignment = align
    if keep_next:
        p.paragraph_format.keep_with_next = True
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def h1(text):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(text)
    r.bold = True
    return p


def h2(text):
    return add_p(text, style="Head -2", keep_next=True)


def h5(text):
    p = doc.add_paragraph(style="Heading 5")
    r = p.add_run(text)
    r.bold = True
    p.paragraph_format.keep_with_next = True
    return p


def bullet(text):
    return add_p(text, style="bullet list")


def figure(path, num, caption, width=3.1):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = add_p(f"{num}-р зураг. {caption}", style="Head 6")
    c.paragraph_format.space_before = Pt(4)
    c.paragraph_format.space_after = Pt(6)


def equation(parts, num):
    """parts: [(text, {i, sub, sup})] — Times New Roman."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    p = doc.add_paragraph(style="equation")
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Pt(126), WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(Pt(252), WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(Inches(1.56), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Inches(3.1), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.add_run("\t")
    for t, f in parts:
        r = p.add_run(t)
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        r._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        r.italic = True if f.get("i") else None
        if f.get("sub"):
            r.font.subscript = True
        if f.get("sup"):
            r.font.superscript = True
    r = p.add_run(f"\t({num})")
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def eq(s):
    """'D|i' тэмдэглэгээ: 'x_{i}' -> доод индекс, italic хувьсагч '*x*'."""
    import re
    out = []
    for tok in re.split(r"(\*[^*]+\*|_\{[^}]+\}|\^\{[^}]+\})", s):
        if not tok:
            continue
        if tok.startswith("*"):
            out.append((tok[1:-1], {"i": True}))
        elif tok.startswith("_{"):
            out.append((tok[2:-1], {"sub": True, "i": True}))
        elif tok.startswith("^{"):
            out.append((tok[2:-1], {"sup": True}))
        else:
            out.append((tok, {}))
    return out


def _cell_shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(num, title, header, rows, widths, notes=(), font=7.5, bold_rows=()):
    t = add_p(title.upper(), style="Table", keep_next=True)
    t.paragraph_format.space_before = Pt(8)
    hp = doc.add_paragraph(style="table head")
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_before = Pt(2)
    hp.paragraph_format.space_after = Pt(3)
    hp.paragraph_format.keep_with_next = True
    pPr = hp._element.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", "0"), ("w:numId", "0")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        numPr.append(e)
    pPr.insert(1, numPr)
    r = hp.add_run(f"{num}-р хүснэгт.")
    r.italic = True
    tb = doc.add_table(rows=len(rows) + 1, cols=len(header))
    tb.alignment = 1
    tblPr = tb._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "BFBFBF")
        borders.append(b)
    tblPr.append(borders)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    cm = OxmlElement("w:tblCellMar")
    for side, v in (("left", "50"), ("right", "50")):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), v)
        e.set(qn("w:type"), "dxa")
        cm.append(e)
    tblPr.append(cm)
    for ri, row in enumerate([header] + rows):
        tr = tb.rows[ri]
        if ri == 0:
            trPr = tr._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
        for ci, val in enumerate(row):
            cell = tr.cells[ci]
            cell.width = Inches(widths[ci])
            p = cell.paragraphs[0]
            p.style = doc.styles["table col head" if ri == 0 else "table copy"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or ci > 0) else WD_ALIGN_PARAGRAPH.LEFT
            for tt, f in parse_inline(str(val)):
                rr = p.add_run(tt)
                rr.font.size = Pt(font)
                if ri == 0 or f.get("b") or (ri - 1) in bold_rows:
                    rr.bold = True
                if f.get("i"):
                    rr.italic = True
                if f.get("sub"):
                    rr.font.subscript = True
            if ri == 0:
                _cell_shade(cell, "9CC2E5")
    for i, w in enumerate(widths):
        for cell in tb.columns[i].cells:
            cell.width = Inches(w)
    grid = tb._element.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    for n in notes:
        np_ = doc.add_paragraph(style="table footnote")
        pPr = np_._element.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        for tag, val in (("w:ilvl", "0"), ("w:numId", "0")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            numPr.append(e)
        pPr.insert(1, numPr)
        np_.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rr = np_.add_run(n)
        rr.italic = True
        rr.font.size = Pt(6)
    sp = doc.add_paragraph(style="Normal")
    sp.paragraph_format.space_after = Pt(2)
    sp.paragraph_format.line_spacing = Pt(4)


def ref(text):
    return add_p(text, style="references")


# ------------------------------------------------------------------ толгой хэсэг
TITLE = ("УХААЛАГ ГАР УТАСНЫ ХЭРЭГЛЭГЧИЙН ӨГӨГДӨЛ БОЛОН ХИЙМЭЛ ОЮУН УХААНД СУУРИЛАН УМАЙН АГШИЛТЫГ "
         "БОДИТ ХУГАЦААНД ҮНЭЛЭХ, ЭРСДЭЛИЙГ АНГИЛАХ СИСТЕМИЙН СУДАЛГАА БА ХӨГЖҮҮЛЭЛТ")
set_runs(els[0], [(TITLE, {})], first_rpr(els[0]))
rp = first_rpr(els[2])
set_runs(els[2], [("[Оюутны овгийн эхний үсэг]. [НЭР]", {"b": True}), ("1", {"sup": True}),
                  (", [Удирдагчийн эрдмийн зэрэг] [овгийн эхний үсэг]. [НЭР]", {"b": True}), ("1", {"sup": True})], rp)
rp = first_rpr(els[3])
set_runs(els[3], [("1", {"sup": True}), ("Монгол Улс, Улаанбаатар, ШУТИС, [сургуулийн бүтэн нэр], [салбарын бүтэн нэр]", {})], rp)
els[4].getparent().remove(els[4])
rp = first_rpr(els[5])
set_runs(els[5], [("Холбоо барих зохиогчийн и-мэйл хаяг: [нэр]@must.edu.mn", {}), ("1", {"sup": True})], rp)

ABSTRACT = (
    "Хураангуй: Жирэмсний сүүлийн үед умайн агшилтын давтамж, үргэлжлэх хугацаа болон хавсарсан шинж тэмдгийг "
    "гэрийн нөхцөлд зөв бүртгэж, эмнэлгийн тусламж хэзээ авахаа шийдэх нь хэрэглэгчид хүндрэлтэй байдаг. Гар утасны "
    "олон аппликейшн агшилтын хугацааг хэмжих боломжтой ч хэрэглэгчийн шинж тэмдэг, чөлөөт бичвэрийг хугацааны "
    "мэдээлэлтэй нэгтгэн шинжилдэггүй. Энэхүү судалгааны зорилго нь ухаалаг гар утсаар бүртгэсэн агшилтын хугацааны "
    "цуваа, өвдөлтийн үнэлгээ, сонгосон шинж тэмдэг болон монгол хэлээрх чөлөөт бичвэрийг нэгтгэн боловсруулж, "
    "анхаарах түвшнийг бага, дунд, өндөр гэж ангилах шийдвэр гаргалтыг дэмжих системийн прототип боловсруулж, "
    "үнэлэхэд оршино. Системд агшилтын цуваанаас хорин шинж чанар гаргах алгоритм, латин үсгийн хөрвүүлэлт, "
    "үгүйсгэлийн илрүүлэлт болон үсгийн алдааг тэсвэрлэх тааруулалт бүхий монгол хэлний байгалийн хэлний "
    "боловсруулалтын модуль, XGBoost ангилагч, аюултай шинж тэмдэг илэрвэл өндөр түвшин оноох хамгаалалтын давхаргыг "
    "хэрэгжүүлэв. Бодит өвчтөний мэдээлэл ашиглалгүйгээр нийтэд ил эмнэлгийн зөвлөмжид тулгуурласан зохиомол өгөгдөл "
    "үүсгэж, хэрэглэгч тус бүрээр хуваасан таван давталттай туршилтаар дүрэмд суурилсан, зөвхөн хугацааны цуваа, "
    f"зөвхөн шинж тэмдэг болон хосолмол загваруудыг харьцуулав. Хосолмол загварын macro-F1 үзүүлэлт {M('HYBRID/XGB','macro_f1'):.3f} "
    f"болж, зөвхөн хугацааны цувааны загвар ({M('TS/XGB','macro_f1'):.3f}) болон дүрэмд суурилсан аргаас "
    f"({M('A2 Rule (timing+red flags)','macro_f1'):.3f}) статистикийн ач холбогдолтой давуу байв. Хамгаалалтын давхарга "
    f"нэмэхэд өндөр түвшний тохиолдлыг илрүүлэх мэдрэг чанар {M('HYBRID/XGB','rec_high'):.3f}-аас "
    f"{M('HYBRID/XGB+guard','rec_high'):.3f} болж өссөн ч хуурамч дохио нэмэгдсэн. Хэлний модуль урьд өмнө хараагүй "
    f"хэллэг дээр {nlpc['B_heldout']['rule_neg_fuzzy']['F1']:.3f} F1 үзүүлсэн нь бодит хэрэглэгчийн бичвэрийн корпус "
    f"шаардлагатайг харуулав. Серверийн хариу өгөх хугацааны медиан {lat['server_p50']:.1f} миллисекунд байв. Судалгааны шинэлэг тал нь агшилтын хугацааны цуваа болон монгол хэлээрх шинж тэмдгийн бичвэрийг нэг "
    "загварт нэгтгэж, аюулгүй байдлын давхаргатай хослуулсанд оршино. Прототип нь мобайл клиент, REST API, өгөгдлийн "
    "сан, автомат тест бүхий бүрэн программ хангамжийн шийдэл юм. Үр дүн нь "
    "эмнэлзүйн үр нөлөөг нотлохгүй бөгөөд эмнэлгийн мэргэжилтний баталгаажуулалт, эмнэлзүйн туршилт шаардлагатай."
)
set_runs(els[8], [(ABSTRACT, {})])
set_runs(els[9], [("Түлхүүр үг: Хугацааны цуваа, байгалийн хэлний боловсруулалт, XGBoost, шийдвэр гаргалтын "
                   "дэмжлэг, mHealth", {})])

# хуучин бичвэр, эмхэтгэлийн нүүрийг устгах; 2 баганатай sectPr-ийг баримтын төгсгөлд шилжүүлэх
two_col = copy.deepcopy(els[91].find(qn("w:pPr")).find(qn("w:sectPr")))
for e in els[10:217]:
    body.remove(e)
old = body.find(qn("w:sectPr"))
body.replace(old, two_col)

# ================================================================== I. УДИРТГАЛ
h1("УДИРТГАЛ")
add_p("Гар утсанд суурилсан эрүүл мэндийн технологи (mHealth) нь эрүүл мэндийн мэдээллийг хувь хүнд ойртуулж буй "
      "бөгөөд Дэлхийн эрүүл мэндийн байгууллага (ДЭМБ) дижитал интервенцийг эрүүл мэндийн тогтолцоог бэхжүүлэх "
      "хэрэгсэл гэж үзэн, түүнийг нотолгоонд тулгуурлан нэвтрүүлэхийг зөвлөсөн [1]. Нөгөө талаар Монгол Улсын "
      "Хүний хувийн мэдээлэл хамгаалах тухай хууль эрүүл мэндийн мэдээллийг эмзэг мэдээлэлд хамааруулж, түүнийг "
      "боловсруулахад тусгай шаардлага тавьдаг [2] тул ийм системийн нууцлалыг дизайны түвшинд шийдэх шаардлагатай.")
add_p("Жирэмсний хугацаанд хэрэглэгчид агшилтын давтамж, үргэлжлэх хугацааг ажиглаж, эмнэлэгт хэзээ хандахаа "
      "шийддэг. Жишээлбэл, 37 долоо хоногоос өмнө цагт зургаа ба түүнээс олон агшилт илрэх, ус гоожих, цус гарах нь "
      "дутуу төрөлтийн шинж байж болох тул эмнэлгийн тусламж авахыг зөвлөдөг [3]. Төрөх хугацаандаа агшилт таван "
      "минут тутам, нэг минут үргэлжилж, нэг цаг тогтвортой байх (5-1-1 дүрэм) нь эмнэлэгт очих нийтлэг заавар юм "
      "[4]. Гэвч өвдөлт, стрессийн үед хугацааг гараар тэмдэглэх, дунджийг тооцох, олон шинж тэмдгийг зэрэг үнэлэх нь "
      "алдаа гаргах эрсдэлтэй.")
add_p("Жирэмсний аппликейшнүүдийн системтэй үнэлгээгээр судалсан 29 аппын 86% нь агшилтын таймер зэрэг тусгай "
      "функцтэй боловч зөвхөн 28% нь шинжлэх ухааны эх сурвалжаас иш татсан байна [5]. Бидний мэдэж буйгаар "
      "хэрэглэгчийн монгол хэлээрх чөлөөт бичвэрийг агшилтын хугацааны цуваатай нэгтгэн, тайлбарлагдахуйц "
      "үнэлгээ өгдөг нээлттэй судалгаа хараахан хийгдээгүй байна.")
add_p("Энэхүү ажлын зорилго нь ухаалаг гар утсаар цуглуулсан агшилтын хугацааны цуваа болон шинж тэмдгийн "
      "мэдээллийг хиймэл оюун ухаан (AI) ашиглан боловсруулж, анхаарах түвшнийг ангилах шийдвэр гаргалтыг дэмжих "
      "(decision-support) программ хангамжийн прототип боловсруулах явдал юм. Систем нь эмчийн оношийг орлохгүй. "
      "Ажлын үндсэн хувь нэмэр:")
bullet("гараар оруулсан алдааг цэвэрлэж, агшилтын цуваанаас 20 шинж чанар гаргах алгоритм;")
bullet("латин-кирилл хөрвүүлэлт, үгүйсгэл, үсгийн алдааг тэсвэрлэх монгол хэлний шинж тэмдэг ялгах NLP модуль;")
bullet("дүрэм, хугацааны цуваа, шинж тэмдэг болон хосолмол загваруудыг ижил нөхцөлд харьцуулах, дахин давтах "
       "боломжтой туршилтын протокол ба зохиомол өгөгдөл үүсгэгч;")
bullet("мобайл клиент, REST API, хамгаалалтын давхарга, баталгаажуулалттай яаралтай холбоо бүхий ажилладаг прототип.")
add_p("Судалгаагаар дараах асуултад хариулна: агшилтын хэв маягийг автоматаар тодорхойлж болох уу (RQ1); чөлөөт "
      "бичвэрийг бүтэцтэй шинж чанар болгож болох уу (RQ2); хосолмол арга нь зөвхөн хугацааны цуваанаас давуу юу "
      "(RQ3); систем бодит хугацаанд хариу өгөх үү (RQ4); аюулгүй байдлын шаардлагыг хэрхэн хангах вэ (RQ5).")

# ================================================================== II. ХОЛБОГДОХ СУДАЛГАА
h1("ХОЛБОГДОХ СУДАЛГАА")
add_p("Умайн агшилтыг объектив хэмжих чиглэлд электрогистерограмм (EHG) дохиогоор дутуу төрөлтийг таамаглах "
      "судалгаа өргөн хийгдсэн. Fele-Žorž нар EHG бичлэгийн нээлттэй сан (TPEHG) бүрдүүлж, шугаман ба шугаман бус "
      "дохио боловсруулах аргуудыг харьцуулсан [6] бол Fergus нар уг сан дээр машин сургалтын ангилагчдыг туршсан "
      "[7]. Эдгээр нь тусгай мэдрэгч шаарддаг тул хэрэглэгчийн гараар оруулсан хугацааны өгөгдөлд шууд "
      "хэрэглэгдэхгүй.")
add_p("Эмнэлзүйн бичвэрээс мэдээлэл ялгах хэрэглээний тоймд дүрэмд суурилсан аргууд практикт давамгайлсаар "
      "байгааг тэмдэглэсэн [8]. Үгүйсгэлийг илрүүлэх NegEx алгоритм [9] нь “цус гараагүй” мэтийн хэллэгийг зөв "
      "тайлбарлахад чухал. Том хэлний загвар (LLM) уян хатан боловч бодит бус агуулга үүсгэх (hallucination) "
      "эрсдэлтэй [10] тул аюулгүй байдалд нөлөөлөх хэрэглээнд тайлбарлагдахуйц, шалгагдахуйц аргыг эхлээд ашиглах нь "
      "зүйтэй. Эмнэлзүйн шийдвэр гаргалтыг дэмжих системийн тоймд хэт олон анхааруулга хэрэглэгчийг ядраах "
      "(alert fatigue) эрсдэлийг онцолсон [11] нь мэдрэг чанар ба хуурамч дохионы тэнцвэрийг судлах үндэслэл болно.")

# ================================================================== III. АРГА ЗҮЙ
h1("СУДАЛГААНЫ АРГА ЗҮЙ")
h2("1. Системийн архитектур")
add_p("Системийн бүтцийг 1-р зурагт үзүүлэв. Гар утасны апп нь агшилт бүрийн эхлэх, дуусах хугацаа, өвдөлтийн "
      "түвшин (1–10), есөн шинж тэмдгийн сонголт, чөлөөт бичвэр болон жирэмсний долоо хоногийг бүртгэнэ. Сүлжээгүй "
      "үед төхөөрөмж дээр агшилтын тооцоо ба аюулгүйн дүрэм ажиллана. Сүлжээтэй үед өгөгдөл REST API-д илгээгдэж, "
      "хугацааны цуваа ба NLP шинж чанарууд нэгтгэгдэн (feature fusion) ангилагчид орно. Хариуд анхаарах түвшин, "
      "түүний тайлбар, зөвлөмж буцна.")
figure(FIG / "fig1_arch.png", 1, "Системийн ерөнхий архитектур")

h2("2. Өгөгдөл ба шошгожилт")
add_p("Бодит жирэмсэн хүний мэдээлэл ашиглах нь ёс зүйн зөвшөөрөл шаарддаг тул энэ үе шатанд зохиомол "
      "(synthetic) өгөгдөл ашиглав. Давталт бүрт 700 хийсвэр хэрэглэгч 1–3 удаагийн, 20–120 минутын бүртгэл хийнэ. "
      "Агшилтын хэв маягийг жигд бус, эхэн үеийн, хилийн болон идэвхтэй гэсэн дөрвөн төрлөөр интервал, хэлбэлзлийн "
      "коэффициент (CV), үргэлжлэх хугацааны мужаар загварчилж, хэрэглэгчдийн 35%-ийг 37 долоо хоногоос өмнөх "
      "хугацаатай болгов. Хэрэглэгчийн алдааг 0–8 секундын хариу үйлдлийн саатал, 4–6 секундын санамсаргүй шуугиан, "
      "0–12% орхигдсон агшилт, 3% санамсаргүй товшилтоор загварчлав. Шинж тэмдгийг сонголтоор 30–80%, чөлөөт "
      "бичвэрээр хэрэглэгчээс хамааран 30–95% магадлалтай мэдээлнэ; бичвэрт үгүйсгэсэн хэллэг, 5% үсгийн алдаа, "
      "латин үсгээр бичих тохиолдол багтана.")
add_p("Шошгыг хэрэглэгчийн оруулсан утгаар бус, жинхэнэ (далд) агшилтын цуваа ба шинж тэмдгээс 1-р хүснэгтийн "
      "дүрмээр тогтоосон. Мэргэжилтнүүдийн санал зөрөлдөөнийг илэрхийлэх зорилгоор шошгын 3%-ийг зэргэлдээ ангилал "
      "руу санамсаргүй сольсон. Иймд туршилт нь загварын эмнэлзүйн үнэн зөвийг бус, шуугиан, дутуу мэдээлэлд тэсвэртэй "
      "эсэхийг шалгана.")
table(1, "Судалгааны шошгоны дүрэм", ["Түвшин", "Нөхцөл (жинхэнэ утгаар)"], [
    ["HIGH", "Аюултай шинж тэмдэг: цус гарах, ус гоожих, ургийн хөдөлгөөн багасах, тасралтгүй хүчтэй өвдөлт, "
             "халуурах, хүчтэй толгой өвдөх эсвэл хараа бүрэлзэх; <37 д.х. үед цагт ≥6 агшилт [3]; ≥37 д.х. үед 5-1-1 "
             "нөхцөл 1 цаг тогтвортой [4]"],
    ["MEDIUM", "<37 д.х.: цагт 4–6 агшилт, эсвэл ≥2 агшилт ба нуруу өвдөх/аарцаг дарагдах; ≥37 д.х.: 5-1-1 нөхцөл "
               "1 цаг хүрээгүй, эсвэл интервал ≤10 мин ба CV ≤0.35"],
    ["LOW", "Дээрх нөхцөл биелээгүй"],
], [0.55, 2.55], notes=["д.х. – жирэмсний долоо хоног. Босго утгууд эмнэлгийн мэргэжилтний баталгаажуулалт шаардана."])

h2("3. Хугацааны цувааны шинж чанар")
add_p("Бүртгэлийг эхлэх хугацаагаар эрэмбэлж, 10 секундээс богино товшилтыг хасаж, 180 секундээс урт бичлэгийг "
      "таслан, давхцсан бичлэгийг нэгтгэнэ. Сүүлийн 60 минутын цонхонд _n_ агшилтын үргэлжлэх хугацаа _D_~i~, "
      "эхлэх хугацаа _s_~i~ бол дундаж үргэлжлэх хугацаа, дундаж интервал, тогтмол байдал ба давтамжийг дараах байдлаар "
      "тооцно:")
equation(eq("*D̄* = (1/*n*) Σ *D*_{i}"), 1)
equation(eq("*Ī* = Σ (*s*_{i+1} − *s*_{i}) / (*n* − 1)"), 2)
equation(eq("*CV*_{I} = σ_{I} / *Ī*,   *R* = 1 / (1 + *CV*_{I})"), 3)
equation(eq("*f* = 60 *n* / max(*T*_{w}, 15)"), 4)
add_p("Энд σ~I~ – интервалын стандарт хазайлт, _T_~w~ – цонхны бодит урт (минут). Эдгээрээс гадна үргэлжлэх "
      "хугацаа ба интервалын хамгийн бага, их утга, стандарт хазайлт, хамгийн бага квадратын аргаар тооцсон чиг "
      "хандлага, интервал ≤5 минут ба үргэлжлэх ≥60 секунд байгаа агшилтын хувь, 5-1-1 нөхцөл тасралтгүй хадгалагдсан "
      "минут, нийт бүртгэлийн хугацаа, жирэмсний долоо хоног болон дутуу хугацааны тэмдэглэгээ зэрэг нийт 20 шинж "
      "чанар гаргав. Өвдөлтөөс 4, сонголтоос 9, NLP-ээс 10, нийт 43 шинж чанар хосолмол загварт орно.")

h2("4. Монгол хэлний NLP модуль")
add_p("Модуль нь бичвэрээс онош бус, есөн шинж тэмдгийн бүтэцтэй тэмдэглэгээ гаргана. Үүнд: (i) латин үсгээр "
      "бичсэн бол кирилл рүү хөрвүүлэх (ts→ц, kh→х г.м.); (ii) латин бичлэгийн хоёрдмол утгыг арилгахаар ү, ө, о "
      "эгшгийг у болгон нугалах; (iii) цэг таслалыг хадгалан токенчлох; (iv) монгол хэл залгамал тул үгийн язгуурын "
      "дарааллаар угтвар тааруулах; (v) дөрөв ба түүнээс урт язгуурт Левенштейний зай ≤1 байхыг зөвшөөрөх; (vi) "
      "NegEx [9]-ийн зарчмаар “-гүй”, “-аагүй” залгавар эсвэл “биш”, “үгүй”, “алга”, “байхгүй” үгийг гурван токены "
      "хүрээнд илрүүлж, цэг, таслал, “гэхдээ”, “харин” үгээр хүрээг таслах алхамтай. “Хөдлөхгүй”, “намдахгүй”, "
      "“тасралтгүй” мэт үг өөрөө шинж тэмдгийг илэрхийлэх тул үгүйсгэлээс хасав. “8/10” хэлбэрийн өвдөлтийн оноог "
      "мөн ялгана. Толь бичгийг эцэслэсний дараа өөр хэллэгтэй 1200 бичвэр бүхий “B” багцыг бичиж, ерөнхийлөх "
      "чадварыг шалгасан. Харьцуулах суурь болгон тэмдэгтийн 2–5-грамм TF-IDF ба логистик регрессийг ашиглав.")

h2("5. Загварууд ба туршилтын дизайн")
add_p("Дараах загваруудыг харьцуулав: A1 – хугацааны дүрэм (1-р хүснэгтийн босгоор); A2 – A1 дээр мэдээлсэн "
      "аюултай шинж тэмдгийг нэмсэн дүрэм; B – зөвхөн хугацааны цуваа; C – зөвхөн шинж тэмдэг, өвдөлт, NLP; D – "
      "бүх шинж чанарыг нэгтгэсэн хосолмол загвар. B ба D-д логистик регресс (LR), санамсаргүй ой (RF) [12], XGBoost "
      "[13]-ийг scikit-learn [14] орчинд тогтмол гиперпараметрээр (XGBoost: 300 мод, гүн 4, сургалтын хурд 0.05, "
      "ангиллын тэнцвэржүүлсэн жин) сургав. D+guard хувилбарт хэрэглэгч аюултай шинж тэмдэг сонгосон эсвэл бичсэн "
      "бол ангилагчийн гаралтаас үл хамааран HIGH оноодог хамгаалалтын давхарга нэмэв.")
add_p("Нэг хэрэглэгчийн бүртгэлүүд хоорондоо хамааралтай тул өгөгдлийг хэрэглэгчээр 70/30 харьцаагаар хувааж, "
      "мэдээлэл алдагдлаас (leakage) сэргийлэв [15]. Өгөгдөл үүсгэх ба хуваалтыг таван санамсаргүй үрээр давтав. "
      "Нарийвчлал, macro-F1, HIGH ангиллын нарийвчлал (precision) ба мэдрэг чанар (recall)-ыг тооцож, загвар хоорондын "
      "ялгааг таван давталтын нийт тестийн дээжид (n=2147) McNemar тестээр шалгав. Ablation туршилтаар хугацааны "
      "цуваан дээр өвдөлт, сонголт, NLP шинж чанарыг дараалан нэмэв.")

h2("6. Нууцлал, аюулгүй байдал, ёс зүй")
add_p("Хэрэглэгчийг нэргүй UUID ба JWT токеноор таньж, серверт зөвхөн үнэлгээний хураангуйг хадгалж, чөлөөт "
      "бичвэрийг хадгалахгүй. Байршлыг зөвхөн хэрэглэгч яаралтай мессеж илгээхийг баталгаажуулсны дараа авна. "
      "103 руу автомат дуудлага хийхгүй бөгөөд хариу бүрт “эмнэлгийн онош биш” гэсэн анхааруулга хавсаргана.")

# ================================================================== IV. ПРОТОТИП
h1("ПРОТОТИПИЙН ХЭРЭГЖҮҮЛЭЛТ")
add_p("Backend-ийг Python 3.11, FastAPI, XGBoost ашиглан хэрэгжүүлж, SQLite өгөгдлийн сантай (PostgreSQL-ээр "
      "солих боломжтой) Docker контейнерт багцлав. API нь /api/v1/auth/anonymous, /api/v1/assess, /health гэсэн "
      "цэгтэй бөгөөд хариунд анхаарах түвшин, ангиллын магадлал, хамгаалалтын давхарга ажилласан эсэх, дүрмийн "
      "суурь үнэлгээ, тайлбар, бичвэрээс танигдсан шинж тэмдэг орно. Android/iOS-д зориулсан Flutter клиентийн эх "
      "кодыг бичиж, туршилт ба үзүүлэнд гар утасны дэлгэцэнд тохирсон вэб клиентийг ашиглав (2-р зураг). Интерфейс нэг гараар ашиглах том товч, "
      "бага текст, өнгө ба үгээр давхар илэрхийлсэн түвшинтэй. Агшилтын тооцоо, NLP, дүрэм, баталгаажуулалт, "
      "эрхийн шалгалтыг хамарсан 21 автомат тест (pytest) бүгд амжилттай давсан.")
figure(FIG / "fig2_ui.png", 2, "Мобайл клиентийн дэлгэц")

# ================================================================== V. ҮР ДҮН
h1("ҮР ДҮН")
e1 = ex["seed0"]["E1"]
h2("1. Агшилтын тооцооны нарийвчлал")
add_p(f"Жинхэнэ цуваатай харьцуулахад цэвэрлэсэн өгөгдлөөс тооцсон дундаж интервалын дундаж үнэмлэхүй алдаа "
      f"(MAE) {e1['clean_iv_mae_min']:.2f} минут (цэвэрлээгүй үед {e1['raw_iv_mae_min']:.2f}), дундаж үргэлжлэх "
      f"хугацааных {e1['clean_d_mae_s']:.2f} секунд (цэвэрлээгүй үед {e1['raw_d_mae_s']:.2f}) байв. Интервалын "
      f"алдааны медиан {e1['clean_iv_median']:.2f} минут бөгөөд том алдаа нь голчлон орхигдсон агшилтаас үүдэн "
      "интервал хоёр дахин уртсахад гарч байна (RQ1).")

h2("2. NLP модулийн үр дүн")
a, b = nlpc["A_test"], nlpc["B_heldout"]
rows = []
for key, name in [("rule_base", "Толь (үгүйсгэлгүй)"), ("rule_neg", "+ Үгүйсгэл"), ("rule_neg_fuzzy", "+ Fuzzy тааруулалт"),
                  ("charLR", "Тэмдэгт n-грамм + LR"), ("rule_or_charLR", "Толь ∪ LR")]:
    rows.append([name, f3(a[key]["P"]), f3(a[key]["R"]), f3(a[key]["F1"]), f3(b[key]["P"]), f3(b[key]["R"]), f3(b[key]["F1"])])
table(2, "Шинж тэмдэг ялгах NLP-ийн үр дүн (micro)", ["Арга", "A: P", "A: R", "A: F1", "B: P", "B: R", "B: F1"], rows,
      [1.2, 0.32, 0.32, 0.32, 0.32, 0.32, 0.32],
      notes=[f"A – сургалтын хэллэгтэй тест (n={a['n']}); B – урьд хараагүй хэллэг (n={b['n']}). P – precision, R – recall."],
      bold_rows=(2,))
add_p(f"2-р хүснэгтээс харахад үгүйсгэлийн илрүүлэлт precision-ийг {a['rule_base']['P']:.2f}-аас "
      f"{a['rule_neg']['P']:.2f} болгож, fuzzy тааруулалт үсгийн алдааг нөхсөн. Гэвч B багц дээр F1 "
      f"{b['rule_neg_fuzzy']['F1']:.3f} болж буурсан нь толь бичиг шинэ хэллэгт бүрэн ерөнхийлөгдөхгүйг харуулна. "
      f"Сургалттай тэмдэгт n-граммын загвар B дээр бүр доогуур ({b['charLR']['F1']:.3f}) байсан бол хоёрыг нэгтгэхэд "
      f"recall {b['rule_or_charLR']['R']:.3f} болж өссөн (RQ2).")

h2("3. Ангиллын гүйцэтгэл")
order = [("A1 Rule (timing)", "A1 Дүрэм"), ("A2 Rule (timing+red flags)", "A2 Дүрэм + шинж"),
         ("TS/LR", "B LR"), ("TS/RF", "B RF"), ("TS/XGB", "B XGB"), ("SYM/XGB", "C XGB"),
         ("HYBRID/LR", "D LR"), ("HYBRID/RF", "D RF"), ("HYBRID/XGB", "D XGB"), ("HYBRID/XGB+guard", "D XGB+guard")]
rows = [[n, f3(M(k, "acc")), pm(k, "macro_f1"), f3(M(k, "prec_high")), f3(M(k, "rec_high"))] for k, n in order]
table(3, "Ангиллын үр дүн (5 давталтын дундаж)", ["Загвар", "Accuracy", "Macro-F1", "HIGH P", "HIGH R"], rows,
      [0.95, 0.52, 0.78, 0.45, 0.45],
      notes=["B – зөвхөн хугацааны цуваа, C – зөвхөн шинж тэмдэг/NLP, D – хосолмол. ± – стандарт хазайлт."],
      bold_rows=(8, 9))
add_p(f"3-р хүснэгт ба 3-р зурагт үзүүлснээр хосолмол XGBoost загвар хамгийн өндөр macro-F1 "
      f"({pm('HYBRID/XGB','macro_f1')}) үзүүлж, зөвхөн хугацааны цувааны XGBoost-оос "
      f"(χ²={T['HYBRID/XGB vs TS/XGB']['stat']:.1f}, p<0.001) болон A2 дүрмээс (χ²={T['HYBRID/XGB vs A2']['stat']:.1f}, "
      f"p<0.001) давуу байв. Зөвхөн хугацааны цувааны загвар ч A1 дүрмээс давуу "
      f"(χ²={T['TS/XGB vs A1']['stat']:.1f}, p<0.001) байсан. HIGH ангиллын хамгийн өндөр мэдрэг чанарыг A2 дүрэм "
      f"({M('A2 Rule (timing+red flags)','rec_high'):.3f}) үзүүлсэн боловч precision нь "
      f"{M('A2 Rule (timing+red flags)','prec_high'):.3f} байв. Хамгаалалтын давхарга бүхий D XGB+guard нь "
      f"{M('HYBRID/XGB+guard','rec_high'):.3f} мэдрэг чанар, {M('HYBRID/XGB+guard','prec_high'):.3f} precision-тэй, "
      f"macro-F1-ээр A2-оос давуу (χ²={T['HYBRID+guard vs A2']['stat']:.1f}, p<0.001) байв (RQ3, RQ5).")
figure(FIG / "fig3_models.png", 3, "Загваруудын macro-F1 ба HIGH ангиллын мэдрэг чанар")
cms = ex["cm_sum"]
import numpy as np
def miss(k):
    c = np.array(cms[k], float)
    return (c[2, 0] + c[2, 1]) / c[2].sum() * 100
def fa(k):
    c = np.array(cms[k], float)
    return c[0, 2] / c[0].sum() * 100
add_p(f"Алдааны матрицаас (4-р зураг) HIGH тохиолдлыг доогуур ангилах алдаа зөвхөн хугацааны цуваанд "
      f"{miss('TS/XGB'):.0f}%, хосолмол загварт {miss('HYBRID/XGB'):.0f}%, хамгаалалтын давхаргатай үед "
      f"{miss('HYBRID/XGB+guard'):.0f}% болж буурсан бол LOW тохиолдлыг HIGH гэх хуурамч дохио {fa('HYBRID/XGB'):.0f}%-"
      f"иас {fa('HYBRID/XGB+guard'):.0f}% болж өссөн. Үлдсэн алдаа нь хэрэглэгч огт мэдээлээгүй шинж тэмдэг ба шошгоны "
      "шуугианаас үүдэлтэй.")
figure(FIG / "fig4_cm.png", 4, "Алдааны матриц (5 давталтын нийлбэр, мөрөөр хувь)")

h2("4. Ablation ба шинж чанарын ач холбогдол")
add_p(f"Хугацааны цуваан дээр ({M('TS/XGB','macro_f1'):.3f}) өвдөлт нэмэхэд {M('TS+PAIN/XGB','macro_f1'):.3f}, "
      f"сонголт нэмэхэд {M('TS+CHECK/XGB','macro_f1'):.3f}, NLP нэмэхэд {M('TS+NLP/XGB','macro_f1'):.3f}, бүгдийг "
      f"нэгтгэхэд {M('HYBRID/XGB','macro_f1'):.3f} болсон. Хамгийн их нэмэрийг чөлөөт бичвэр өгсөн бөгөөд NLP-тэй "
      f"хувилбар ба бүрэн загварын ялгаа ач холбогдолгүй (p={T['HYBRID/XGB vs TS+NLP/XGB']['p']:.3f}). 5-р зурагт "
      "дутуу хугацааны тэмдэглэгээ, цагт ногдох агшилт, дундаж интервал, 5-1-1 тогтвортой хугацаа тэргүүлж, "
      "NLP-ээр илэрсэн ургийн хөдөлгөөн багасах, ус гоожих шинж дараагийн байранд орсон.")
figure(FIG / "fig5_imp.png", 5, "Хосолмол загварын шинж чанарын ач холбогдол (эхний 12)")

h2("5. Хариу өгөх хугацаа")
table(4, "Системийн хариу өгөх хугацаа (мс)", ["Нөхцөл", "p50", "p95", "p99"], [
    ["Серверийн боловсруулалт", f"{lat['server_p50']:.1f}", f"{lat['server_p95']:.1f}", "–"],
    ["Дараалсан 1000 хүсэлт (RTT)", f"{lat['seq_rtt_p50']:.1f}", f"{lat['seq_rtt_p95']:.1f}", f"{lat['seq_rtt_p99']:.1f}"],
    ["20 зэрэгцээ хэрэглэгч (RTT)", f"{lat['conc_p50']:.1f}", f"{lat['conc_p95']:.1f}", "–"],
], [1.6, 0.5, 0.5, 0.5],
    notes=[f"2 vCPU, localhost, нэг хүсэлтэд дунджаар {lat['n_contr_mean']:.1f} агшилт. Зэрэгцээ үед "
           f"{lat['conc_throughput_rps']:.0f} хүсэлт/с, алдаа {lat['conc_errors']}."])
add_p("4-р хүснэгтээс харахад нэг үнэлгээний серверийн боловсруулалт хэдхэн миллисекунд үргэлжилж, 20 зэрэгцээ "
      "хэрэглэгчтэй үед ч 95-р перцентиль 0.3 секундээс бага байсан нь бодит хугацааны шаардлагыг хангана (RQ4). "
      "Гар утасны сүлжээний саатлыг энэ хэмжилт хамруулаагүй.")

# ================================================================== VI. ХЭЛЭЛЦҮҮЛЭГ
h1("ХЭЛЭЛЦҮҮЛЭГ")
add_p("Үр дүн нь хугацааны цуваан дээр шинж тэмдгийн мэдээллийг нэмэх нь ангиллыг сайжруулдаг гэсэн таамаглалыг "
      "дэмжив. Гэхдээ шошго нь бидний тодорхойлсон дүрмээс гаралтай тул загвар эдгээр дүрмийг шуугиантай өгөгдлөөс "
      "дахин сурч буй хэрэг бөгөөд тоон утгууд нь эмнэлзүйн гүйцэтгэл биш. Мөн хэрэглэгч шинж тэмдгийг хэр олон "
      "мэдээлэх магадлалыг таамгаар тогтоосон тул ablation-ы харьцаа бодит хэрэглээнд өөрчлөгдөж болно.")
add_p("Аюулгүй байдлын хувьд HIGH тохиолдлыг алгасах нь хуурамч дохионоос илүү хор хөнөөлтэй тул прототипт "
      "хамгаалалтын давхаргыг анхдагч болгосон. Гэвч хуурамч дохио нэмэгдэх нь хэрэглэгчийг ядраах эрсдэлтэй [11] тул "
      "бодит хэрэглээнд эмнэлгийн мэргэжилтэнтэй хамтран тэнцвэрийг тогтоох шаардлагатай. NLP-ийн ерөнхийлөх "
      "чадварын зөрүү (A ба B) нь зөвшөөрөлтэйгөөр цуглуулж, тэмдэглэгээ хийсэн монгол хэлний бодит корпус "
      "хамгийн чухал дараагийн алхам болохыг харуулна. LLM ашиглах тохиолдолд схемийн шалгалт ба дүрмийн "
      "баталгаажуулалттай хослуулах нь зүйтэй [10].")
add_p("Хязгаарлалт: өгөгдөл зохиомол; ухаалаг утас агшилтыг физиологийн мэдрэгчээр хэмжихгүй; босго утгуудыг "
      "эмнэлгийн мэргэжилтэн баталгаажуулаагүй; хариу өгөх хугацааг локал сүлжээнд хэмжсэн. Иймд эмнэлзүйн "
      "хэрэглээнээс өмнө ёс зүйн зөвшөөрөлтэй проспектив судалгаа заавал шаардлагатай.")

# ================================================================== ДҮГНЭЛТ
h5("ДҮГНЭЛТ")
add_p("Энэхүү ажлаар ухаалаг гар утсаар бүртгэсэн агшилтын хугацааны цуваа, өвдөлт, шинж тэмдэг болон монгол "
      "хэлээрх чөлөөт бичвэрийг нэгтгэн анхаарах түвшнийг бодит хугацаанд ангилах шийдвэр гаргалтыг дэмжих "
      "системийн ажилладаг прототипийг боловсруулж, дахин давтах боломжтой туршилтаар үнэлэв. Зохиомол өгөгдөл "
      f"дээр хосолмол загвар macro-F1 {M('HYBRID/XGB','macro_f1'):.3f} үзүүлж, зөвхөн хугацааны цуваа болон "
      "дүрэмд суурилсан аргаас статистикийн ач холбогдолтой давуу байсан бөгөөд хамгаалалтын давхарга HIGH "
      f"тохиолдлын мэдрэг чанарыг {M('HYBRID/XGB+guard','rec_high'):.3f} хүртэл нэмэгдүүлэв. Монгол хэлний NLP "
      "модуль үгүйсгэл, латин бичлэг, үсгийн алдааг зохицуулж чадсан ч шинэ хэллэгт ерөнхийлөх чадвар хязгаарлагдмал "
      "байв. Цаашид бодит, зөвшөөрөлтэй өгөгдөл цуглуулах, эмнэлгийн мэргэжилтэнтэй хамтран шошгоны дүрмийг "
      "баталгаажуулах, ухаалаг цаг зэрэг мэдрэгч нэгтгэх, монгол хэлний эрүүл мэндийн NLP загвар сургах, "
      "эмнэлзүйн туршилт хийхээр төлөвлөж байна.")

# ================================================================== НОМ ЗҮЙ
h5("АШИГЛАСАН МАТЕРИАЛ, НОМ ЗҮЙ")
refs = [
    "World Health Organization, WHO guideline: Recommendations on digital interventions for health system "
    "strengthening. Geneva: WHO, 2019.",
    "Монгол Улсын Их Хурал, “Хүний хувийн мэдээлэл хамгаалах тухай хууль,” 2021 оны 12-р сарын 17. [Онлайн]. "
    "Хаяг: https://legalinfo.mn/mn/detail?lawId=16390288615991",
    "Eunice Kennedy Shriver National Institute of Child Health and Human Development, “What are the symptoms of "
    "preterm labor?” 2023. [Online]. Available: https://www.nichd.nih.gov/health/topics/preterm/conditioninfo/symptoms",
    "Lamaze International, “Stages of labor.” [Online]. Available: https://lamaze.org/stages-of-labor",
    "G. Frid, K. Bogaert, and K. T. Chen, “Mobile health apps for pregnant women: Systematic search, evaluation, "
    "and analysis of features,” J. Med. Internet Res., vol. 23, no. 10, e25667, 2021.",
    "G. Fele-Žorž, G. Kavšek, Ž. Novak-Antolič, and F. Jager, “A comparison of various linear and non-linear "
    "signal processing techniques to separate uterine EMG records of term and pre-term delivery groups,” Med. Biol. "
    "Eng. Comput., vol. 46, no. 9, pp. 911–922, 2008.",
    "P. Fergus, P. Cheung, A. Hussain, D. Al-Jumeily, C. Dobbins, and S. Iram, “Prediction of preterm deliveries "
    "from EHG signals using machine learning,” PLoS ONE, vol. 8, no. 10, e77154, 2013.",
    "Y. Wang, L. Wang, M. Rastegar-Mojarad, S. Moon, F. Shen, N. Afzal, S. Liu, Y. Zeng, S. Mehrabi, S. Sohn, and "
    "H. Liu, “Clinical information extraction applications: A literature review,” J. Biomed. Inform., vol. 77, "
    "pp. 34–49, 2018.",
    "W. W. Chapman, W. Bridewell, P. Hanbury, G. F. Cooper, and B. G. Buchanan, “A simple algorithm for "
    "identifying negated findings and diseases in discharge summaries,” J. Biomed. Inform., vol. 34, no. 5, "
    "pp. 301–310, 2001.",
    "Z. Ji, N. Lee, R. Frieske, T. Yu, D. Su, Y. Xu, E. Ishii, Y. J. Bang, A. Madotto, and P. Fung, “Survey of "
    "hallucination in natural language generation,” ACM Comput. Surv., vol. 55, no. 12, art. 248, 2023.",
    "R. T. Sutton, D. Pincock, D. C. Baumgart, D. C. Sadowski, R. N. Fedorak, and K. I. Kroeker, “An overview of "
    "clinical decision support systems: Benefits, risks, and strategies for success,” npj Digit. Med., vol. 3, "
    "art. 17, 2020.",
    "L. Breiman, “Random forests,” Mach. Learn., vol. 45, no. 1, pp. 5–32, 2001.",
    "T. Chen and C. Guestrin, “XGBoost: A scalable tree boosting system,” in Proc. 22nd ACM SIGKDD Int. Conf. "
    "Knowledge Discovery and Data Mining, 2016, pp. 785–794.",
    "F. Pedregosa, G. Varoquaux, A. Gramfort, V. Michel, B. Thirion, O. Grisel, M. Blondel, P. Prettenhofer, "
    "R. Weiss, V. Dubourg, J. Vanderplas, A. Passos, D. Cournapeau, M. Brucher, M. Perrot, and E. Duchesnay, "
    "“Scikit-learn: Machine learning in Python,” J. Mach. Learn. Res., vol. 12, pp. 2825–2830, 2011.",
    "S. Kaufman, S. Rosset, C. Perlich, and O. Stitelman, “Leakage in data mining: Formulation, detection, and "
    "avoidance,” ACM Trans. Knowl. Discov. Data, vol. 6, no. 4, art. 15, 2012.",
]
for r_ in refs:
    ref(r_)

# ---- OOXML схемийн дарааллыг засах
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
             "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
             "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
TBL_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
             "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
             "tblCaption", "tblDescription"]


def reorder(el, order):
    kids = list(el)
    key = lambda k: order.index(k.tag.split("}")[1]) if k.tag.split("}")[1] in order else len(order)
    for k in kids:
        el.remove(k)
    for k in sorted(kids, key=key):
        el.append(k)


for ppr in body.iter(qn("w:pPr")):
    reorder(ppr, PPR_ORDER)
for tp in body.iter(qn("w:tblPr")):
    reorder(tp, TBL_ORDER)

doc.core_properties.title = "Умайн агшилтыг бодит хугацаанд үнэлэх AI систем"
doc.core_properties.author = ""
out = W / "paper" / "Author_School_MM-YYYY_Umain_agshilt_AI.docx"
doc.save(str(out))
print(out)
